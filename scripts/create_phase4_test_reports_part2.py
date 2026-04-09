"""
Phase 4 Testing Reports Generator - Part 2
Creates individual DOCX reports for Maintenance Mode, Branch Deploys, and DNS tests
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

# Brand colors
PRIMARY_COLOR = RGBColor(151, 192, 232)
PRIMARY_DARK = RGBColor(23, 64, 105)

def set_heading_format(run, size=14, color=PRIMARY_DARK, bold=True):
    """Apply consistent heading formatting"""
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = 'Inter'

def add_status_badge(paragraph, status):
    """Add colored status badge"""
    run = paragraph.add_run(f' [{status}]')
    run.font.bold = True
    if status == 'PASSED':
        run.font.color.rgb = RGBColor(0, 128, 0)
    elif status == 'FAILED':
        run.font.color.rgb = RGBColor(255, 0, 0)

def create_maintenance_mode_report():
    """Create Maintenance Mode Testing Report"""
    doc = Document()

    title = doc.add_heading('Maintenance Mode Testing Report', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_heading_format(title.runs[0], size=18, color=PRIMARY_DARK)

    doc.add_paragraph()
    table = doc.add_table(rows=8, cols=2)
    table.style = 'Light Grid Accent 1'

    metadata = [
        ('Test Category', 'Maintenance Mode Testing'),
        ('Test Date', 'February 11, 2026'),
        ('Tester', 'Keenan Husselmann'),
        ('Environment', 'Production (https://x4o.co.za)'),
        ('Total Test Cases', '6'),
        ('Passed', '6'),
        ('Failed', '0'),
        ('Pass Rate', '100%')
    ]

    for i, (label, value) in enumerate(metadata):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
        table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True

    doc.add_heading('Test Objective', 1)
    doc.add_paragraph(
        'Verify maintenance mode can be activated and displays correctly while preserving '
        'access to static assets. Ensure the maintenance page is accessible for testing.'
    )

    doc.add_heading('Test Summary', 1)

    summary_table = doc.add_table(rows=7, cols=3)
    summary_table.style = 'Medium Grid 1 Accent 1'

    headers = ['Test Case', 'Status', 'Notes']
    for i, header in enumerate(headers):
        summary_table.rows[0].cells[i].text = header
        summary_table.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True

    summary_data = [
        ('TC-MAINT-001: Page Exists', '[OK] PASSED', 'Accessible via direct URL'),
        ('TC-MAINT-002: Content Complete', '[OK] PASSED', 'All elements present'),
        ('TC-MAINT-003: Responsive', '[OK] PASSED', 'Mobile-friendly'),
        ('TC-MAINT-004: Asset Preservation', '[OK] PASSED', 'Images accessible'),
        ('TC-MAINT-005: Activation Process', '[OK] PASSED', 'Simple toggle'),
        ('TC-MAINT-006: Direct URL Test', '[OK] PASSED', 'Preview without activating')
    ]

    for i, (tc, status, notes) in enumerate(summary_data, start=1):
        summary_table.rows[i].cells[0].text = tc
        summary_table.rows[i].cells[1].text = status
        summary_table.rows[i].cells[2].text = notes

    doc.add_paragraph()
    conclusion = doc.add_paragraph('Overall Test Result: ')
    run = conclusion.add_run('[OK] 6/6 PASSED (100%)')
    run.font.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 128, 0)

    output_path = 'docs/phase4-testing/test-report-03-maintenance-mode.docx'
    doc.save(output_path)
    print(f'[OK] Created: {output_path}')

def create_branch_deploy_report():
    """Create Branch Deploy Separation Testing Report"""
    doc = Document()

    title = doc.add_heading('Branch Deploy Separation Testing Report', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_heading_format(title.runs[0], size=18, color=PRIMARY_DARK)

    doc.add_paragraph()
    table = doc.add_table(rows=8, cols=2)
    table.style = 'Light Grid Accent 1'

    metadata = [
        ('Test Category', 'Branch Deploy Separation Testing'),
        ('Test Date', 'February 11, 2026'),
        ('Tester', 'Keenan Husselmann'),
        ('Environment', 'Netlify (main & staging branches)'),
        ('Total Test Cases', '7'),
        ('Passed', '7'),
        ('Failed', '0'),
        ('Pass Rate', '100%')
    ]

    for i, (label, value) in enumerate(metadata):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
        table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True

    doc.add_heading('Test Objective', 1)
    doc.add_paragraph(
        'Verify staging and production branches deploy separately without cross-contamination. '
        'Ensure proper Git workflow separation.'
    )

    doc.add_heading('Test Summary', 1)

    summary_table = doc.add_table(rows=8, cols=3)
    summary_table.style = 'Medium Grid 1 Accent 1'

    headers = ['Test Case', 'Status', 'Notes']
    for i, header in enumerate(headers):
        summary_table.rows[0].cells[i].text = header
        summary_table.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True

    summary_data = [
        ('TC-BRANCH-001: Configuration', '[OK] PASSED', 'Contexts in netlify.toml'),
        ('TC-BRANCH-002: Production Branch', '[OK] PASSED', 'Deploys to x4o.co.za'),
        ('TC-BRANCH-003: Staging Branch', '[OK] PASSED', 'Deploys to Netlify subdomain'),
        ('TC-BRANCH-004: Git Workflow', '[OK] PASSED', 'Documented and enforced'),
        ('TC-BRANCH-005: Independence', '[OK] PASSED', 'No cross-contamination'),
        ('TC-BRANCH-006: Maintenance Separation', '[OK] PASSED', 'Can diff per branch'),
        ('TC-BRANCH-007: PR Previews', '[OK] PASSED', 'Auto-generated URLs')
    ]

    for i, (tc, status, notes) in enumerate(summary_data, start=1):
        summary_table.rows[i].cells[0].text = tc
        summary_table.rows[i].cells[1].text = status
        summary_table.rows[i].cells[2].text = notes

    doc.add_paragraph()
    conclusion = doc.add_paragraph('Overall Test Result: ')
    run = conclusion.add_run('[OK] 7/7 PASSED (100%)')
    run.font.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 128, 0)

    output_path = 'docs/phase4-testing/test-report-04-branch-deploys.docx'
    doc.save(output_path)
    print(f'[OK] Created: {output_path}')

def create_dns_resolution_report():
    """Create DNS Resolution Testing Report"""
    doc = Document()

    title = doc.add_heading('DNS Resolution Testing Report', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_heading_format(title.runs[0], size=18, color=PRIMARY_DARK)

    doc.add_paragraph()
    table = doc.add_table(rows=8, cols=2)
    table.style = 'Light Grid Accent 1'

    metadata = [
        ('Test Category', 'DNS Resolution Testing'),
        ('Test Date', 'February 11, 2026'),
        ('Tester', 'Keenan Husselmann'),
        ('Environment', 'Production (x4o.co.za)'),
        ('Total Test Cases', '7'),
        ('Passed', '7'),
        ('Failed', '0'),
        ('Pass Rate', '100%')
    ]

    for i, (label, value) in enumerate(metadata):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
        table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True

    doc.add_heading('Test Objective', 1)
    doc.add_paragraph(
        'Verify domain resolves correctly to Netlify infrastructure, redirects function properly, '
        'HTTPS is enforced, and email DNS records remain unchanged.'
    )

    doc.add_heading('Test Summary', 1)

    summary_table = doc.add_table(rows=8, cols=3)
    summary_table.style = 'Medium Grid 1 Accent 1'

    headers = ['Test Case', 'Status', 'Notes']
    for i, header in enumerate(headers):
        summary_table.rows[0].cells[i].text = header
        summary_table.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True

    summary_data = [
        ('TC-DNS-001: Root Domain (A)', '[OK] PASSED', 'Points to Netlify IP'),
        ('TC-DNS-002: WWW Subdomain (CNAME)', '[OK] PASSED', 'Points to Netlify site'),
        ('TC-DNS-003: WWW to Non-WWW', '[OK] PASSED', '301 redirect configured'),
        ('TC-DNS-004: HTTPS Enforcement', '[OK] PASSED', 'Auto SSL via Lets Encrypt'),
        ('TC-DNS-005: Email MX Preserved', '[OK] PASSED', 'Google Workspace active'),
        ('TC-DNS-006: Custom 404', '[OK] PASSED', 'Branded error page'),
        ('TC-DNS-007: Propagation Time', '[OK] PASSED', 'Fully propagated')
    ]

    for i, (tc, status, notes) in enumerate(summary_data, start=1):
        summary_table.rows[i].cells[0].text = tc
        summary_table.rows[i].cells[1].text = status
        summary_table.rows[i].cells[2].text = notes

    doc.add_paragraph()
    conclusion = doc.add_paragraph('Overall Test Result: ')
    run = conclusion.add_run('[OK] 7/7 PASSED (100%)')
    run.font.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 128, 0)

    output_path = 'docs/phase4-testing/test-report-05-dns-resolution.docx'
    doc.save(output_path)
    print(f'[OK] Created: {output_path}')

if __name__ == '__main__':
    print('Generating Phase 4 Test Reports (Part 2)...\n')
    create_maintenance_mode_report()
    create_branch_deploy_report()
    create_dns_resolution_report()
    print('\n[OK] Part 2 Complete: 3 reports generated')
