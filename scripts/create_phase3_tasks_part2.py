"""
Phase 3 Individual Task Documentation Generator - Part 2
Service Pages Documentation
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level=1):
    """Add a styled heading"""
    heading = doc.add_heading(text, level=level)
    if level == 1:
        heading.runs[0].font.color.rgb = RGBColor(23, 64, 105)
        heading.runs[0].font.size = Pt(24)
    elif level == 2:
        heading.runs[0].font.color.rgb = RGBColor(23, 64, 105)
        heading.runs[0].font.size = Pt(18)
    return heading

def add_screenshot_section(doc, title, description):
    """Add a screenshot placeholder section"""
    doc.add_heading(f"Screenshot: {title}", level=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("[INSERT SCREENSHOT HERE]")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(151, 192, 232)
    p = doc.add_paragraph()
    p.add_run("Screenshot Description: ").bold = True
    p.add_run(description)
    p.space_after = Pt(12)

def add_table(doc, data, has_header=True):
    """Add a formatted table"""
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.style = 'Light Grid Accent 1'
    for i, row_data in enumerate(data):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            cell = row.cells[j]
            cell.text = str(cell_text)
            if has_header and i == 0:
                cell.paragraphs[0].runs[0].font.bold = True
                cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(23, 64, 105)
    return table

def add_technical_definitions(doc, definitions):
    """Add technical jargon definitions section"""
    doc.add_page_break()
    add_heading(doc, "Technical Jargon Definitions", level=1)
    for term, definition in definitions.items():
        p = doc.add_paragraph()
        p.add_run(f"{term}: ").bold = True
        p.add_run(definition)

# ============================================================
# TASK 3: Build Consulting & Advisory Page
# ============================================================
def create_task_consulting_page():
    doc = Document()

    add_heading(doc, 'Task: Build Consulting & Advisory Page', level=1)
    add_heading(doc, 'consulting-and-advisory-services.astro', level=2)

    meta_data = [
        ['Task ID', 'TASK-013 (Development Phase)'],
        ['File', 'src/pages/consulting-and-advisory-services.astro'],
        ['Status', '✅ Complete'],
        ['Lines of Code', '~180 lines']
    ]
    add_table(doc, meta_data, has_header=False)

    doc.add_paragraph()
    doc.add_paragraph('─' * 60)
    doc.add_paragraph()

    add_heading(doc, 'Overview', level=2)
    doc.add_paragraph(
        'This page showcases X4O\'s core consulting and advisory services across three main categories: '
        'Project Management, Advisory Services, and Renewable Energy Solutions. The page uses glassmorphic '
        'card designs with detailed service descriptions and credentials.'
    )

    doc.add_paragraph()

    add_heading(doc, 'Page Structure', level=2)
    doc.add_paragraph(
        '1. Hero Banner\n'
        '   • Header image: x4o_banner.jpg\n'
        '   • Page title: "Consulting & Advisory Services"\n'
        '   • Breadcrumb navigation\n\n'
        '2. Introduction Section\n'
        '   • Company expertise statement\n'
        '   • Value proposition\n'
        '   • Years of experience\n\n'
        '3. Service Cards (3 Main Services)\n'
        '   • Project Management\n'
        '     - Full project lifecycle management\n'
        '     - Risk assessment and mitigation\n'
        '     - Stakeholder coordination\n'
        '     - PMI/PRINCE2 methodologies\n'
        '   • Advisory Services\n'
        '     - Strategic business planning\n'
        '     - Process optimization\n'
        '     - Compliance consulting\n'
        '     - Regulatory guidance\n'
        '   • Renewable Energy Solutions\n'
        '     - Solar energy implementation\n'
        '     - Energy efficiency audits\n'
        '     - Sustainability consulting\n'
        '     - Grid integration\n\n'
        '4. Call-to-Action Section\n'
        '   • "Ready to get started?" heading\n'
        '   • Contact button\n'
        '   • Email and phone info',
        style='No Spacing'
    )

    doc.add_paragraph()
    doc.add_paragraph('─' * 60)
    doc.add_paragraph()

    add_screenshot_section(
        doc,
        'Full Page Desktop View',
        'Full desktop screenshot showing hero banner with x4o_banner.jpg at top, page title, '
        'introduction paragraph, and 3 service cards in horizontal layout. Each card shows icon, '
        'title, bullet points, and glass effect styling.'
    )

    doc.add_paragraph()

    add_screenshot_section(
        doc,
        'Service Cards Grid',
        'Close-up of the 3 service cards showing: Project Management (left) with project icon, '
        'Advisory Services (center) with lightbulb icon, Renewable Energy (right) with sun icon. '
        'Cards have glassmorphic background, subtle borders, and hover effects.'
    )

    doc.add_paragraph()

    add_screenshot_section(
        doc,
        'Mobile Stacked Layout',
        'Mobile view showing service cards stacked vertically in single column, full-width, '
        'with proper spacing. Banner image responsive, text readable, and CTA button full-width.'
    )

    doc.add_paragraph()
    doc.add_paragraph('─' * 60)
    doc.add_paragraph()

    add_heading(doc, 'Content Details', level=2)

    services_table = [
        ['Service', 'Key Features', 'Target Clients'],
        ['Project Management', 'Lifecycle management, Risk mitigation, Stakeholder coordination', 'Enterprises, Government, NGOs'],
        ['Advisory Services', 'Strategic planning, Process optimization, Compliance', 'SMBs, Corporates, Public sector'],
        ['Renewable Energy', 'Solar implementation, Energy audits, Sustainability', 'Industrial, Commercial, Residential']
    ]
    add_table(doc, services_table)

    doc.add_paragraph()

    add_heading(doc, 'SEO Configuration', level=2)
    seo_table = [
        ['Meta Tag', 'Value'],
        ['Title', 'Consulting & Advisory Services | X4O Consultants'],
        ['Description', 'Expert project management, business advisory, and renewable energy solutions in South Africa'],
        ['Keywords', 'consulting, advisory, project management, renewable energy, South Africa']
    ]
    add_table(doc, seo_table)

    doc.add_paragraph()

    add_heading(doc, 'SDLC Best Practices Applied', level=2)
    doc.add_paragraph(
        '✅ Content Hierarchy: Clear information architecture\n'
        '✅ Responsive Design: Mobile-first breakpoints\n'
        '✅ SEO Optimization: Unique meta tags\n'
        '✅ User Experience: Clear CTAs and navigation\n'
        '✅ Accessibility: Semantic HTML, proper headings\n'
        '✅ Performance: Optimized images, static generation\n'
        '✅ Code Reusability: Shared Layout component\n'
        '✅ Maintainability: Consistent styling patterns',
        style='No Spacing'
    )

    definitions = {
        'Project Management': 'Discipline of planning, organizing, and managing resources to achieve specific goals',
        'Advisory Services': 'Professional guidance and recommendations for business strategy and operations',
        'Renewable Energy': 'Energy from sources that naturally replenish (solar, wind, hydro)',
        'Risk Mitigation': 'Strategies to reduce probability or impact of potential threats',
        'Stakeholder Coordination': 'Managing communication and expectations of project participants',
        'Compliance Consulting': 'Guidance on meeting regulatory and legal requirements',
        'Energy Audit': 'Assessment of energy consumption to identify efficiency improvements',
        'Glassmorphism': 'Modern UI design featuring frosted-glass effects with backdrop blur',
        'Breadcrumb Navigation': 'Secondary navigation showing user\'s location in site hierarchy',
        'CTA (Call-to-Action)': 'Element designed to prompt immediate user response'
    }
    add_technical_definitions(doc, definitions)

    return doc

# ============================================================
# TASK 4: Build Coaching Services Page
# ============================================================
def create_task_coaching_page():
    doc = Document()

    add_heading(doc, 'Task: Build Coaching Services Page', level=1)
    add_heading(doc, 'coaching-services.astro', level=2)

    meta_data = [
        ['Task ID', 'TASK-014 (Development Phase)'],
        ['File', 'src/pages/coaching-services.astro'],
        ['Status', '✅ Complete'],
        ['Lines of Code', '~220 lines']
    ]
    add_table(doc, meta_data, has_header=False)

    doc.add_paragraph()
    doc.add_paragraph('─' * 60)
    doc.add_paragraph()

    add_heading(doc, 'Overview', level=2)
    doc.add_paragraph(
        'The coaching services page presents X4O\'s four specialized coaching programs: Professional '
        'Development, Business Growth, Project Management, and Executive Coaching. Each program features '
        'detailed descriptions, target audiences, and outcomes in glassmorphic card layouts.'
    )

    doc.add_paragraph()

    add_heading(doc, 'Page Structure', level=2)
    doc.add_paragraph(
        '1. Hero Section\n'
        '   • Gradient background\n'
        '   • Page title: "Coaching Services"\n'
        '   • Tagline: "Transform Your Career and Business"\n\n'
        '2. Introduction\n'
        '   • Coaching philosophy statement\n'
        '   • Approach methodology\n'
        '   • Success metrics\n\n'
        '3. Coaching Programs (4 Cards)\n'
        '   • Professional Development Coaching\n'
        '     - Career advancement strategies\n'
        '     - Leadership skill building\n'
        '     - Performance improvement\n'
        '   • Business Growth Coaching\n'
        '     - Strategic planning facilitation\n'
        '     - Market expansion guidance\n'
        '     - Operational efficiency\n'
        '   • Project Management Coaching\n'
        '     - PM certification preparation\n'
        '     - Methodology implementation\n'
        '     - Team leadership training\n'
        '   • Executive Coaching\n'
        '     - C-suite leadership development\n'
        '     - Decision-making frameworks\n'
        '     - Organizational transformation\n\n'
        '4. Booking CTA\n'
        '   • "Ready to start your journey?" section\n'
        '   • Link to Book Coaching Sessions page\n'
        '   • Contact information',
        style='No Spacing'
    )

    doc.add_paragraph()
    doc.add_paragraph('─' * 60)
    doc.add_paragraph()

    add_screenshot_section(
        doc,
        '4 Coaching Cards Grid (Desktop)',
        'Desktop view showing all 4 coaching program cards in 2x2 grid layout. Each card displays: '
        'unique icon, program title in bold, descriptive paragraph, bullet points for key features, '
        'and glassmorphic styling with hover effects.'
    )

    doc.add_paragraph()

    add_screenshot_section(
        doc,
        'Single Coaching Card Detail',
        'Close-up of Professional Development Coaching card showing: circular gradient icon with '
        'person symbol, "Professional Development Coaching" heading, description text, bullet list '
        'with 3-4 key features, glass-card background with subtle border and shadow.'
    )

    doc.add_paragraph()

    add_screenshot_section(
        doc,
        'Mobile View (Stacked Cards)',
        'Mobile screenshot showing coaching cards stacked vertically, each card full-width, '
        'proper spacing between cards, readable text size, and touch-friendly layout.'
    )

    doc.add_paragraph()

    add_screenshot_section(
        doc,
        'Booking CTA Section',
        'Bottom section showing "Ready to start your journey?" heading, descriptive text about '
        'booking process, prominent "Book a Coaching Session" button with gradient styling, '
        'and alternative contact methods below.'
    )

    doc.add_paragraph()
    doc.add_paragraph('─' * 60)
    doc.add_paragraph()

    add_heading(doc, 'Coaching Programs Comparison', level=2)

    programs_table = [
        ['Program', 'Target Audience', 'Duration', 'Outcomes'],
        ['Professional Development', 'Individual contributors, managers', '3-6 months', 'Career advancement, skill mastery'],
        ['Business Growth', 'Entrepreneurs, business owners', '6-12 months', 'Revenue growth, market expansion'],
        ['Project Management', 'PM professionals, team leads', '3-4 months', 'Certification, methodology expertise'],
        ['Executive Coaching', 'C-suite, senior leaders', '6-12 months', 'Strategic leadership, org transformation']
    ]
    add_table(doc, programs_table)

    doc.add_paragraph()

    add_heading(doc, 'Card Design Specifications', level=2)
    design_table = [
        ['Element', 'Specification'],
        ['Layout', '2x2 grid desktop, 2x1 tablet, 1x1 mobile'],
        ['Card Background', 'glass-card class (backdrop blur, semi-transparent)'],
        ['Icons', 'Gradient circles with Heroicons SVG'],
        ['Typography', 'Title: text-xl font-bold, Body: text-gray-600'],
        ['Spacing', 'p-8 padding, gap-8 grid gap'],
        ['Hover Effect', 'hover:-translate-y-1 (lift), transition-transform']
    ]
    add_table(doc, design_table)

    doc.add_paragraph()

    add_heading(doc, 'SDLC Best Practices Applied', level=2)
    doc.add_paragraph(
        '✅ Information Architecture: Clear program categorization\n'
        '✅ User Journey: Direct path to booking page\n'
        '✅ Visual Consistency: Matching design patterns from other pages\n'
        '✅ Responsive Grid: 2x2 → 2x1 → 1x1 breakpoints\n'
        '✅ Accessibility: Descriptive headings, semantic structure\n'
        '✅ SEO: Unique content for each coaching type\n'
        '✅ Performance: Minimal JavaScript, static generation\n'
        '✅ Maintainability: Reusable card pattern',
        style='No Spacing'
    )

    definitions = {
        'Coaching': 'Structured professional relationship focused on goal achievement and personal growth',
        'Professional Development': 'Process of improving skills and competencies for career advancement',
        'Executive Coaching': 'Leadership development for senior executives and C-suite leaders',
        'Business Growth': 'Strategies and methods to increase revenue and expand market presence',
        'Project Management Certification': 'Formal credentials (PMP, PRINCE2) validating PM expertise',
        'Grid Layout': 'CSS layout system arranging content in rows and columns',
        'Responsive Breakpoints': 'Screen width thresholds where layout changes (mobile → tablet → desktop)',
        'Glassmorphic Card': 'UI element with frosted-glass visual effect and transparency',
        'Hover Effect': 'Visual change when user moves cursor over interactive element',
        'Call-to-Action (CTA)': 'Button or link designed to prompt specific user action'
    }
    add_technical_definitions(doc, definitions)

    return doc

# Save documents
print("Creating Part 2: Service Pages Documentation...")
print("=" * 60)

import os
os.makedirs('docs/phase3-tasks', exist_ok=True)

# Generate Task 3
task3_doc = create_task_consulting_page()
task3_doc.save('docs/phase3-tasks/task-013-build-consulting-advisory-page.docx')
print("Created: task-013-build-consulting-advisory-page.docx")

# Generate Task 4
task4_doc = create_task_coaching_page()
task4_doc.save('docs/phase3-tasks/task-014-build-coaching-services-page.docx')
print("Created: task-014-build-coaching-services-page.docx")

print("=" * 60)
print("Part 2 of task documentation complete...")
print("Continuing with remaining tasks...")
